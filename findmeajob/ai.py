"""Two-stage LLM layer: cheap screen over everything, rich draft for the top few.

Cost lives here, so the two stages are deliberately lopsided:

  screen  batch ~8 jobs per call, JD truncated to ~1400 chars, cheapest model
  draft   one call per job, ~6000 chars of JD, best model, only for the top ~5

Both stages take an optional (provider, model) pair so tests can inject a stub
and so you can point screening at Groq while drafting stays on Claude.
"""
from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from .sources import Job
from .backends import LLMError, Provider, resolve

_FENCE_OPEN = re.compile(r"^\s*```(?:json|JSON)?\s*", re.M)
_FENCE_CLOSE = re.compile(r"\s*```\s*$", re.M)

DRAFT_KEYS = ("fit_summary", "tailored_bullets", "gaps", "cover_note", "questions_to_ask")

# Output ceilings per stage. These are deliberately generous: reasoning models
# (Gemini 2.5+, and anything with thinking on) spend output tokens before the
# answer starts, so a ceiling sized to the visible answer gets consumed and you
# get truncated JSON instead of a result.
SCREEN_MAX_TOKENS = 4000
DRAFT_MAX_TOKENS = 8000
PROFILE_MAX_TOKENS = 4000
PROFILE_RETRY_MAX_TOKENS = 6000
CUSTOMIZE_MAX_TOKENS = 10000

_ROLE_SECTION = re.compile(
    r"\b(responsibilit(?:y|ies)|what you(?:'|’)ll do|duties|key tasks|role overview|about the role)\b",
    re.I,
)
_REQUIREMENT_SECTION = re.compile(
    r"\b(requirements?|qualifications?|what we(?:'|’)re looking for|skills?)\b",
    re.I,
)


def _screen_excerpt(description: str, limit: int) -> str:
    """Prioritize responsibilities and requirements over employer marketing copy."""
    if not description:
        return ""
    limit = max(200, int(limit))
    blocks = [block.strip() for block in re.split(r"\n+", description) if block.strip()]
    if not blocks:
        return description[:limit]

    role_blocks = [block for block in blocks if _ROLE_SECTION.search(block)]
    requirement_blocks = [block for block in blocks if _REQUIREMENT_SECTION.search(block)]
    priority = []
    for block in role_blocks + requirement_blocks:
        if block not in priority:
            priority.append(block)

    # Keep nearby bullets after section headings; ATS text often puts each list
    # item on its own line, so this captures duties without copying the whole JD.
    for index, block in enumerate(blocks):
        if _ROLE_SECTION.search(block) or _REQUIREMENT_SECTION.search(block):
            for nearby in blocks[index + 1:index + 9]:
                if nearby not in priority:
                    priority.append(nearby)

    context = blocks[:3]
    ordered = []
    for block in context + priority + blocks:
        if block not in ordered:
            ordered.append(block)
    excerpt = "\n".join(ordered)
    return excerpt[:limit]


def extract_pdf_text(pdf: bytes) -> str:
    """Extract text locally for providers that cannot accept PDF documents."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise LLMError("PDF support needs pypdf. Run: pip install pypdf") from exc
    try:
        text = "\n\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages
        ).strip()
    except Exception as exc:  # malformed or encrypted PDF
        raise LLMError(f"could not read PDF text: {exc}") from exc
    if not text:
        raise LLMError(
            "this PDF has no selectable text. Use a text-based PDF or choose "
            "Gemini/Anthropic for scanned PDFs."
        )
    return text


def extract_docx_text(docx: bytes) -> str:
    """Extract paragraphs and table cells locally from a DOCX resume."""
    try:
        from docx import Document
    except ImportError as exc:
        raise LLMError(
            "DOCX support needs python-docx. Run: pip install python-docx"
        ) from exc
    try:
        document = Document(BytesIO(docx))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts).strip()
    except Exception as exc:
        raise LLMError(f"could not read DOCX text: {exc}") from exc
    if not text:
        raise LLMError("this DOCX file contains no readable text")
    return text


def parse_json(raw: str) -> Any:
    """Parse a model reply that is *supposed* to be JSON.

    Models wrap JSON in ```fences```, open with "Here is the JSON:", or return
    an object where you asked for an array. All three show up in practice, so
    this is deliberately forgiving: strip fences, try straight, then fall back
    to the outermost bracketed span.
    """
    if raw is None:
        raise ValueError("empty model reply")
    cleaned = _FENCE_CLOSE.sub("", _FENCE_OPEN.sub("", raw)).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        pass
    # Preamble before the payload, or trailing commentary after it.
    candidates = []
    for opener, closer in (("[", "]"), ("{", "}")):
        i, k = cleaned.find(opener), cleaned.rfind(closer)
        if i != -1 and k > i:
            candidates.append((i, cleaned[i:k + 1]))
    for _, blob in sorted(candidates):
        try:
            return json.loads(blob)
        except json.JSONDecodeError:
            continue
    raise ValueError(f"could not parse JSON from model reply: {cleaned[:300]!r}")


def _as_list(payload: Any) -> list[dict]:
    """Accept [ {...} ], { "jobs": [...] }, or a bare { ... }."""
    if isinstance(payload, list):
        return [p for p in payload if isinstance(p, dict)]
    if isinstance(payload, dict):
        for key in ("jobs", "results", "scores", "items"):
            inner = payload.get(key)
            if isinstance(inner, list):
                return [p for p in inner if isinstance(p, dict)]
        return [payload]
    raise ValueError(f"expected a JSON array of results, got {type(payload).__name__}")


# ---------------------------------------------------------------- profile ---

PROFILE_PROMPT = """Extract a structured job-search profile from this resume.

Return ONLY a JSON object, no prose, no markdown fences:
{
  "name": str,
  "current_title": str,
  "years_experience": number,
  "core_skills": [str],        // 10-20, most load-bearing first
  "domains": [str],            // e.g. "distributed systems", "CDN", "frontend"
  "notable_projects": [str],   // one line each, with impact if stated
  "education": str,
    "target_titles": [str],      // all realistic role families this person can target; include every strong fit, not just one title
    "preferred_locations": [str], // cities/countries named as preferred locations
  "seniority": str             // intern | new-grad | junior | mid | senior | staff
}"""


def build_profile(resume_bytes: bytes | None = None, resume_text: str | None = None,
                  is_pdf: bool = False, provider: Provider | None = None,
                  model: str | None = None) -> dict:
    """Resume (PDF or text) -> profile.json. Uses the draft-stage model."""
    if provider is None or model is None:
        provider, model = resolve("draft")

    if is_pdf and resume_bytes and provider.name not in {"anthropic", "gemini"}:
        raw = provider.complete(
            model, "", f"{PROFILE_PROMPT}\n\n--- RESUME ---\n"
            f"{extract_pdf_text(resume_bytes)}", PROFILE_MAX_TOKENS, json_mode=True)
    elif is_pdf and resume_bytes:
        try:
            raw = provider.complete_document(
                model, PROFILE_PROMPT, resume_bytes, PROFILE_MAX_TOKENS)
        except LLMError as e:
            raise LLMError(
                f"{e}\nTip: export your resume to .txt and re-run, or set "
                f"DRAFT_PROVIDER=anthropic|gemini for PDF support."
            ) from e
    else:
        raw = provider.complete(
            model, "", f"{PROFILE_PROMPT}\n\n--- RESUME ---\n{resume_text or ''}",
            PROFILE_MAX_TOKENS, json_mode=True)

    try:
        profile = parse_json(raw)
    except ValueError:
        compact_prompt = """Return ONLY compact valid JSON. No markdown or explanation.
Use at most 8 skills, 5 domains, 5 projects, and 8 target_titles.
Schema: {"name":"", "current_title":"", "years_experience":0,
"core_skills":[], "domains":[], "notable_projects":[], "education":"",
"target_titles":[], "preferred_locations":[], "seniority":""}

RESUME:
"""
        if is_pdf and resume_bytes:
            retry_text = extract_pdf_text(resume_bytes)
        else:
            retry_text = resume_text or ""
        raw = provider.complete(
            model, "", compact_prompt + retry_text,
            PROFILE_RETRY_MAX_TOKENS, json_mode=True)
        profile = parse_json(raw)
    if not isinstance(profile, dict):
        raise ValueError("profile extraction did not return a JSON object")
    return profile


# ----------------------------------------------------------------- screen ---

SCREEN_SYSTEM = """You screen job postings for one candidate. You are strict.

Score 0-10 on genuine fit:
  9-10  strong match, candidate clears the bar and the role is a step up
  7-8   good match, worth applying
  5-6   plausible but real gaps
  0-4   wrong seniority, wrong stack, or a hard requirement the candidate lacks

Seniority mismatch is the most common failure: a 3-year engineer scoring an 8
on a Staff role is wrong. Penalise it hard, in both directions — a senior
engineer does not want an internship either. Do the same for hard requirements
the candidate plainly does not meet: security clearance, a specific degree, a
named technology with a year count they cannot hit, or a country they cannot
work in.

Do not inflate scores to be encouraging. Most postings are a 4.

Return ONLY a JSON array, one object per job, no prose:
[{"job_id": str, "score": number, "reason": str}]
Echo `job_id` back exactly as given. `reason` is one sentence, max 20 words,
concrete about the deciding factor."""


def screen(jobs: list[Job], profile: dict, batch_size: int = 8, jd_chars: int = 1400,
           provider: Provider | None = None, model: str | None = None) -> list[Job]:
    """Stage 1: score every surviving job. Mutates and returns `jobs`.

    A batch that fails to parse logs a warning and is skipped — one bad reply
    must not take down the whole run.
    """
    if provider is None or model is None:
        provider, model = resolve("screen")
    batch_size = max(1, int(batch_size))
    profile_blob = json.dumps(profile, ensure_ascii=False)

    for start in range(0, len(jobs), batch_size):
        batch = jobs[start:start + batch_size]
        payload = [{
            "job_id": j.job_id,
            "company": j.company,
            "title": j.title,
            "location": j.location,
            "description": _screen_excerpt(j.description, jd_chars),
        } for j in batch]

        n = start // batch_size + 1
        try:
            raw = provider.complete(
                model, SCREEN_SYSTEM,
                f"CANDIDATE PROFILE:\n{profile_blob}\n\n"
                f"JOBS:\n{json.dumps(payload, ensure_ascii=False)}",
                SCREEN_MAX_TOKENS, json_mode=True,
            )
            results = {}
            for r in _as_list(parse_json(raw)):
                jid = r.get("job_id")
                if jid:
                    results[str(jid)] = r
        except (LLMError, ValueError, KeyError, TypeError) as e:
            print(f"  ! screen batch {n} failed ({type(e).__name__}: {e}) — skipping")
            continue

        for j in batch:
            r = results.get(j.job_id)
            if not r:
                continue
            try:
                j.score = max(0.0, min(10.0, float(r.get("score", 0))))
            except (TypeError, ValueError):
                j.score = 0.0
            j.reason = str(r.get("reason", "")).strip()

        print(f"  screened {min(start + batch_size, len(jobs))}/{len(jobs)}")

    return jobs


# ------------------------------------------------------------------ draft ---

DRAFT_SYSTEM = """You prepare an application kit for one job.

Hard rule: never invent experience. Every claim must trace to something in the
candidate profile. If the profile does not support a claim, it goes in `gaps`,
not in a bullet.

Return ONLY a JSON object, no prose:
{
  "fit_summary": str,          // 2 sentences: why this is worth their time
  "tailored_bullets": [str],   // 3-4 resume bullets rewritten for THIS job,
                               // using only real experience from the profile,
                               // each with a concrete artefact or number
  "gaps": [str],               // 1-3 honest gaps, each with how to address it
  "cover_note": str,           // 120-160 words. Plain. No "I am writing to
                               // express my interest", no "I am excited to",
                               // no flattery about the company's mission.
                               // Open with a concrete reason they fit this role.
  "questions_to_ask": [str]    // 2 sharp questions that show they read the JD
}"""


def draft(jobs: list[Job], profile: dict, jd_chars: int = 6000,
          provider: Provider | None = None, model: str | None = None) -> list[Job]:
    """Stage 2: full kit for the shortlist. One call per job, best model."""
    if provider is None or model is None:
        provider, model = resolve("draft")
    profile_blob = json.dumps(profile, ensure_ascii=False)

    for j in jobs:
        try:
            raw = provider.complete(
                model, DRAFT_SYSTEM,
                f"CANDIDATE PROFILE:\n{profile_blob}\n\n"
                f"JOB: {j.title} at {j.company} ({j.location or 'location not stated'})\n"
                f"URL: {j.url}\n\n{j.description[:jd_chars]}",
                DRAFT_MAX_TOKENS, json_mode=True,
            )
            kit = parse_json(raw)
            if not isinstance(kit, dict):
                raise ValueError("draft did not return a JSON object")
            # Normalise so the digest template never has to guess.
            j.draft = {
                "fit_summary": str(kit.get("fit_summary") or ""),
                "tailored_bullets": [str(b) for b in (kit.get("tailored_bullets") or [])],
                "gaps": [str(g) for g in (kit.get("gaps") or [])],
                "cover_note": str(kit.get("cover_note") or ""),
                "questions_to_ask": [str(q) for q in (kit.get("questions_to_ask") or [])],
            }
            print(f"  drafted {j.title} @ {j.company}")
        except (LLMError, ValueError, KeyError, TypeError) as e:
            print(f"  ! draft failed for {j.job_id} ({type(e).__name__}: {e})")
            j.draft = {k: ("" if k in ("fit_summary", "cover_note") else []) for k in DRAFT_KEYS}

    return jobs


# ------------------------------------------------------ resume customizer ---

CUSTOMIZE_IDEAL_PROMPT = """Create an ideal resume outline for this job description.

This is a TARGET resume, not a claim about any candidate. Infer the keywords,
responsibilities, evidence and structure that would make a qualified applicant
clear to a recruiter. Do not invent a person's employers, dates or metrics.
Return ONLY JSON:
{
  "target_summary": str,
  "target_skills": [str],
  "target_experience_themes": [str],
  "target_resume": str
}

JOB DESCRIPTION:
"""

CUSTOMIZE_COMPARE_PROMPT = """Compare a candidate's actual resume with an ideal target for a job.

Return ONLY JSON:
{
  "positioning_summary": str,
  "keep": [str],
  "rewrite_suggestions": [str],
  "gaps": [str],
  "tailored_resume": str,
  "truth_check": [str]
}

Rules:
- Never invent experience, employers, dates, skills, education or metrics.
- Only rewrite or reorder evidence present in the actual resume.
- Put unsupported target requirements in gaps, not in the tailored resume.
- Keep the tailored resume practical and ready for the candidate to edit.

IDEAL TARGET:
{ideal}

ACTUAL RESUME:
{actual}
"""


def customize_resume(job: Job, resume_text: str | None = None,
                     resume_bytes: bytes | None = None, is_pdf: bool = False,
                     provider: Provider | None = None,
                     model: str | None = None) -> dict:
    """Build a JD-only target, then compare it with the candidate's resume."""
    if provider is None or model is None:
        provider, model = resolve("draft")
    ideal_raw = provider.complete(
        model, "", CUSTOMIZE_IDEAL_PROMPT + job.description[:12000],
        CUSTOMIZE_MAX_TOKENS, json_mode=True)
    ideal = parse_json(ideal_raw)
    if not isinstance(ideal, dict):
        raise ValueError("ideal resume response was not a JSON object")
    ideal_blob = json.dumps(ideal, ensure_ascii=False)
    compare_prompt = CUSTOMIZE_COMPARE_PROMPT.replace("{ideal}", ideal_blob).replace(
        "{actual}", resume_text or "")
    if is_pdf and resume_bytes and provider.name not in {"anthropic", "gemini"}:
        compare_raw = provider.complete(
            model, "", compare_prompt.replace("{actual}", extract_pdf_text(resume_bytes)),
            CUSTOMIZE_MAX_TOKENS, json_mode=True)
    elif is_pdf and resume_bytes:
        compare_raw = provider.complete_document(
            model, compare_prompt, resume_bytes, CUSTOMIZE_MAX_TOKENS)
    else:
        compare_raw = provider.complete(
            model, "", compare_prompt, CUSTOMIZE_MAX_TOKENS, json_mode=True)
    comparison = parse_json(compare_raw)
    if not isinstance(comparison, dict):
        raise ValueError("resume comparison response was not a JSON object")
    return {"ideal": ideal, "comparison": comparison}


# ------------------------------------------------- offline scorer (no API) ---

def keyword_screen(jobs: list[Job], profile: dict, **_) -> list[Job]:
    """DEV ONLY. Token-overlap stand-in so the pipeline runs with no API key.

    This is not a matcher. It cannot tell a Staff role from a new-grad one and
    it has no idea what the words mean. It exists so `--mock --scorer keyword`
    proves the plumbing end-to-end on a laptop with no secrets configured.
    Never ship a digest built from these scores.
    """
    skills = {s.lower() for s in profile.get("core_skills", []) if s}
    titles = [t.lower() for t in profile.get("target_titles", []) if t]
    for j in jobs:
        blob = f"{j.title} {j.description}".lower()
        hits = sorted(s for s in skills if s in blob)
        overlap = len(hits) / max(len(skills), 1)
        title_bonus = 2.5 if any(t in j.title.lower() for t in titles) else 0.0
        j.score = round(min(10.0, overlap * 12 + title_bonus), 1)
        j.reason = ("[keyword stub] matched: " + ", ".join(hits[:5])) if hits \
            else "[keyword stub] no skill overlap"
    return jobs
